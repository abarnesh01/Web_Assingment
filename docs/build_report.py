from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
source = (ROOT / 'assignment.md').read_text(encoding='utf-8').splitlines()
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.65); section.bottom_margin = Inches(.65)
section.left_margin = Inches(.75); section.right_margin = Inches(.75)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10)
styles['Normal'].font.color.rgb = RGBColor(44, 55, 60)
for name, size, color in [('Title', 30, '142B2A'), ('Heading 1', 19, '183B36'), ('Heading 2', 14, '2F7664'), ('Heading 3', 12, '2F7664')]:
    style = styles[name]; style.font.name = 'Aptos Display'; style.font.size = Pt(size)
    style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(12); style.paragraph_format.space_after = Pt(6)
header = section.header.paragraphs[0]
header.text = 'RELAY  /  REAL-TIME TEAM CHAT'; header.runs[0].font.color.rgb = RGBColor(47,118,100)
header.runs[0].font.bold = True; header.runs[0].font.size = Pt(8)
footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Web Technologies Assignment  •  7 October 2026').font.size = Pt(8)

def shade(cell, color):
    props = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); props.append(shd)
def add_table(lines):
    parsed = [[c.strip() for c in x.strip().strip('|').split('|')] for x in lines]
    rows = [r for r in parsed if not all(re.fullmatch(r':?-{3,}:?', c or '-') for c in r)]
    if not rows: return
    t = doc.add_table(rows=1, cols=len(rows[0])); t.style='Table Grid'
    for i, value in enumerate(rows[0]):
        cell=t.rows[0].cells[i]; cell.text=value; shade(cell,'183B36')
        for run in cell.paragraphs[0].runs:
            run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(8)
    for ri,row in enumerate(rows[1:]):
        cells=t.add_row().cells
        for i,value in enumerate(row):
            cells[i].text=value
            if ri%2==0: shade(cells[i],'F0F5F2')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for run in p.runs: run.font.size=Pt(8)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def add_text(text):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.12
    if text.startswith('> '):
        p.paragraph_format.left_indent=Inches(.18); r=p.add_run(text[2:]); r.italic=True; r.font.color.rgb=RGBColor(47,118,100)
    elif text.startswith('**') and '**' in text[2:]:
        end=text.find('**',2); p.add_run(text[2:end]).bold=True; p.add_run(text[end+2:])
    else:
        p.add_run(text)

idx=0
while idx < len(source):
    line=source[idx].strip()
    if not line: idx+=1; continue
    if line.startswith('![Relay'):
        if (ROOT/'screenshot.png').exists():
            doc.add_picture(str(ROOT/'screenshot.png'), width=Inches(6.7))
            doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        idx+=1; continue
    if line.startswith('*Figure 1.'):
        p=doc.add_paragraph(line.strip('*')); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic=True; p.runs[0].font.size=Pt(8); p.runs[0].font.color.rgb=RGBColor(100,114,113)
        idx+=1; continue
    if line.startswith('|'):
        group=[]
        while idx<len(source) and source[idx].strip().startswith('|'):
            group.append(source[idx].strip()); idx+=1
        add_table(group); continue
    if line.startswith('```'):
        idx+=1; code=[]
        while idx<len(source) and not source[idx].strip().startswith('```'):
            code.append(source[idx]); idx+=1
        idx+=1; p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.15)
        run=p.add_run('\n'.join(code)); run.font.name='Consolas'; run.font.size=Pt(8.5); run.font.color.rgb=RGBColor(38,103,86)
        continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],style='Title'); p.paragraph_format.space_after=Pt(3)
    elif line.startswith('## '): doc.add_heading(line[3:],level=1)
    elif line.startswith('### '): doc.add_heading(line[4:],level=2)
    elif line.startswith('- '):
        p=doc.add_paragraph(line[2:],style='List Bullet'); p.paragraph_format.space_after=Pt(2)
    else: add_text(line)
    idx+=1
out=ROOT/'assignment.docx'; doc.save(out); print(out)
